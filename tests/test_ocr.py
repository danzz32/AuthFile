import os
import unittest
from PIL import Image, ImageDraw

from objects.ai import OcrProcessor

# Caminhos para arquivos de teste
TEST_PDF_PATH = 'authfile/tests/test_document.pdf'
TEST_DOCX_PATH = 'authfile/tests/test_document.docx'
TEST_IMAGE_PATH = 'authfile/tests/test_image.png'


class TestOcrProcessor(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        # Criar arquivos de teste antes dos testes
        # Criar PDF de teste
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.cell(200, 10, txt="Texto de teste para PDF.", ln=True)
        pdf.output(TEST_PDF_PATH)

        # Criar DOCX de teste
        from docx import Document
        doc = Document()
        doc.add_paragraph("Texto de teste para DOCX.")
        doc.save(TEST_DOCX_PATH)

        # Criar imagem de teste
        img = Image.new('RGB', (100, 100), color='white')
        d = ImageDraw.Draw(img)
        d.text((10, 10), "Texto de teste para imagem", fill=(0, 0, 0))
        img.save(TEST_IMAGE_PATH)

    @classmethod
    def tearDownClass(cls):
        # Remover arquivos de teste após os testes
        for file_path in [TEST_PDF_PATH, TEST_DOCX_PATH, TEST_IMAGE_PATH]:
            if os.path.exists(file_path):
                os.remove(file_path)

    def setUp(self):
        self.ocr_processor = OcrProcessor(language='por')

    def test_ocr_process_pdf(self):
        text = self.ocr_processor.process(TEST_PDF_PATH)
        self.assertIn("Texto de teste para PDF.", text, "Texto extraído do PDF está incorreto")

    def test_ocr_process_docx(self):
        text = self.ocr_processor.process(TEST_DOCX_PATH)
        self.assertIn("Texto de teste para DOCX.", text, "Texto extraído do DOCX está incorreto")

    def test_ocr_process_image(self):
        text = self.ocr_processor.process(TEST_IMAGE_PATH)
        self.assertIn("Texto de teste para imagem", text, "Texto extraído da imagem está incorreto")


if __name__ == '__main__':
    unittest.main()
